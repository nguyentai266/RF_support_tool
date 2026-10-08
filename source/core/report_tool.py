
from core.parser import ParserLog

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils import get_column_letter

parser=ParserLog()
def __safe_to_numeric(col):
    try:
        
        return pd.to_numeric(col)
    except (ValueError, TypeError):
       
        return col
    
def export_excel_file(filename, groups, mode):
    regular_font = Font(name='Calibri', size=11, bold=False)
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    yellow_fill = PatternFill(start_color="FFFC0E", end_color="FFFC0E", fill_type="solid")
    red_font = Font(name='Calibri', size=11, color='9C0006', bold=False)
    summary = []
    with pd.ExcelWriter(filename, engine='openpyxl') as w:
        count = 0
        for key, group in groups:
            count += 1
            group = group.apply(__safe_to_numeric)

            print(f"Loading {key}")
            
            if mode == 'station':
                sheet_title = f"Sheet{str(count)}"
            else:
                sheet_title = str(key)

            for char in r"\/?:*[]":
                sheet_title = sheet_title.replace(char, "")
            sheet_title = sheet_title[:31]
            if not sheet_title:
                sheet_title = f"Sheet_{count}"

            mean = group.mean(numeric_only=True)
            result = {'Measurement': filename}
            for col, val in mean.items():
                result[col] = val
            summary.append(result)
            
            df = group.dropna(axis=1, how='any').reset_index(drop=True).T
            
            if mode == 'loop':
                df.columns = [f"LOOP {i}" for i in range(1, len(df.columns) + 1)]
                if 'dut_id' in df.index:
                    new_index_order = ['dut_id'] + [item for item in df.index if item != 'dut_id']
                    df = df.reindex(new_index_order)
            elif mode == "correl":
                if 'station_id' in df.index:
                    df.columns = df.loc['station_id'].astype(str).values
                    df = df.drop('station_id')
            elif mode == "station":
                if 'dut_id' in df.index:
                    df.columns = df.loc['dut_id'].astype(str).values
                    df = df.drop('dut_id')

            df.index.name = "Measurement"
            df.to_excel(w, sheet_name=sheet_title, index=True)
            ws = w.sheets[sheet_title]  # type: ignore
            
           
            ws.views.sheetView[0].showGridLines = True
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.border = Border()
            
            if mode == 'loop':
                last_col_num = ws.max_column
                last_data_col_letter = get_column_letter(last_col_num)
                avg_col = last_col_num + 1
                std_col = last_col_num + 2
                gap_col = last_col_num + 3
                
                ws.cell(row=1, column=avg_col, value="Average")
                ws.cell(row=1, column=std_col, value="Stdev")
                ws.cell(row=1, column=gap_col, value="Gap")
                
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=avg_col, value=f"=AVERAGE(B{r}:{last_data_col_letter}{r})").number_format = '0.0000'
                    ws.cell(row=r, column=std_col, value=f"=STDEV(B{r}:{last_data_col_letter}{r})").number_format = '0.0000'
                    ws.cell(row=r, column=gap_col, value=f"=MAX(B{r}:{last_data_col_letter}{r})-MIN(B{r}:{last_data_col_letter}{r})").number_format = '0.0000'
                gap_letter = get_column_letter(gap_col)
                rule = CellIsRule(operator='greaterThanOrEqual', formula=['1'], fill=red_fill, font=red_font)
                ws.conditional_formatting.add(f"{gap_letter}2:{gap_letter}{ws.max_row}", rule)
                    
            elif mode in ['correl', 'station']: 
                last_data_col_num = ws.max_column  
                last_data_col_letter = get_column_letter(last_data_col_num) 
                
                
                avg_col = last_data_col_num + 1  
                avg_col_letter = get_column_letter(avg_col)
                ws.cell(row=1, column=avg_col, value="Average")
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=avg_col, value=f"=AVERAGE(B{r}:{last_data_col_letter}{r})").number_format = '0.0000'

                
                current_col = avg_col + 1  
                for col_idx in range(2, last_data_col_num + 1):
                    orig_header = ws.cell(row=1, column=col_idx).value  
                    orig_col_letter = get_column_letter(col_idx)        
                    
                    
                    ws.cell(row=1, column=current_col, value=orig_header)
                    
                    for r in range(2, ws.max_row + 1):
                        ws.cell(row=r, column=current_col, value=f"=({orig_col_letter}{r}-${avg_col_letter}{r})^2").number_format = '0.0000'
                    current_col += 1 
                    
                
                gap_col = current_col  
                ws.cell(row=1, column=gap_col, value="Gap")
                gap_letter = get_column_letter(gap_col)
                rule = CellIsRule(operator='greaterThanOrEqual', formula=['1'], fill=red_fill, font=red_font)
                ws.conditional_formatting.add(f"{gap_letter}2:{gap_letter}{ws.max_row}", rule)

                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=gap_col, value=f"=MAX(B{r}:{last_data_col_letter}{r})-MIN(B{r}:{last_data_col_letter}{r})").number_format = '0.0000'

            for r in range(1, ws.max_row + 1):
                ws.cell(row=r, column=avg_col).fill = yellow_fill #type: ignore
            for col in range(1, ws.max_column + 1):
                ws.cell(row=1, column=col).font = regular_font
            for row in range(1, ws.max_row + 1):
                ws.cell(row=row, column=1).font = regular_font
                
        '''if mode == 'station':
            if summary:
                df_sum = pd.DataFrame(summary)
                df_sum.T.to_excel(w, sheet_name='Summary', index=True)
                w.sheets['Summary'].views.sheetView[0].showGridLines = True'''

def crr_and_loop(path,filename,mode):
    limit,summary=parser.summary_data(path)
    summary.drop(columns=['log_path'],inplace=True)
    if mode == "loop":
        groups=parser.group_data(dataFrame=summary,groupBy='dut_id')
        export_excel_file(filename=filename,groups=groups,mode=mode)
    elif mode == "correl":
        groups=parser.group_data(dataFrame=summary,groupBy='dut_id')
        export_excel_file(filename=filename,groups=groups,mode=mode)        
    elif mode == "station":
        groups=parser.group_data(dataFrame=summary,groupBy='station_id')
        export_excel_file(filename=filename,groups=groups,mode=mode) #type

def grr(path,filename):
    limit,summary=parser.summary_data(path)
    grr_df= parser.grr_summary(limit=limit,summary=summary)
    grr_df.drop(columns=['log_path'],inplace=True)
    grr_df.to_csv(filename, index=False)



if __name__ == "__main__":
    path="LOOP"
    loop_path= "LOOP"
    grr_path= "C:/Users/V1531673/Desktop/CELL-04/GRR"
    crr_and_loop(path=path,filename="hhdd",mode="loop")
    #grr(path=grr_path,filename="grr_out")
